"""端到端管线测试:Word(docx)→ PDF → 比对。"""
import io
from pathlib import Path

from docx import Document  # type: ignore[import-untyped]

try:
    import pymupdf as fitz
except ImportError:
    import fitz  # type: ignore

from document_comparison.embed.mock import MockEmbedding
from document_comparison.models import (
    Block,
    Diff,
    PageMeta,
    PageRegion,
    TamperReport,
)
from document_comparison.observability import llm_call_collector
from document_comparison.parsing.pdf import count_pages, extract_text_blocks
from document_comparison.parsing.word_render import render_docx_to_pdf
from document_comparison.pipeline import _parse_source, run_pipeline
from document_comparison.report.builder import burn_pdf


def test_burn_pdf_includes_pdf_added_clause(tmp_path: Path):
    pdf_path = tmp_path / "source.pdf"
    out_path = tmp_path / "annotated.pdf"
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    doc.save(pdf_path)
    doc.close()
    report = TamperReport(
        source="source.docx",
        target=str(pdf_path),
        overall_risk="changed",
        change_status="changed",
        page_meta=[PageMeta(
            page_index=0,
            width_px=1654,
            height_px=2339,
            pdf_width_pt=595,
            pdf_height_pt=842,
        )],
        unmatched_clauses=[Diff(
            alignment_id="al-added",
            status="added",
            page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
            segments=[],
        )],
    )

    burn_pdf(pdf_path, report, out_path)

    annotated = fitz.open(out_path)
    try:
        assert len(list(annotated[0].annots() or [])) == 1
    finally:
        annotated.close()


def test_burn_pdf_renders_deleted_placeholder_as_dashed_rect(tmp_path: Path):
    """deleted 推断占位框(placeholder)用 draw_rect 虚线绘制,而非 add_rect_annot。

    验证:占位 region 不会产生 annot(实线高亮才用 annot),
    而是通过 page.get_drawings() 产生 rect 绘制内容。
    """
    pdf_path = tmp_path / "source.pdf"
    out_path = tmp_path / "annotated.pdf"
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    doc.save(pdf_path)
    doc.close()
    report = TamperReport(
        source="source.docx",
        target=str(pdf_path),
        overall_risk="changed",
        change_status="changed",
        page_meta=[PageMeta(
            page_index=0, width_px=1654, height_px=2339,
            pdf_width_pt=595, pdf_height_pt=842,
        )],
        unmatched_clauses=[Diff(
            alignment_id="al-deleted",
            status="deleted",
            number="第二条",
            page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.2, 0.9, 0.3], kind="placeholder")],
            segments=[],
        )],
    )

    burn_pdf(pdf_path, report, out_path)

    annotated = fitz.open(out_path)
    try:
        page = annotated[0]
        # 占位框走 draw_rect(虚线),不应产生 annot
        annots = list(page.annots() or [])
        assert len(annots) == 0, "placeholder 不应产生 annot(应走 draw_rect 虚线)"
        # draw_rect 会留下 rect 类型的绘制内容
        rect_drawings = [d for d in page.get_drawings() if d.get("rect")]
        assert len(rect_drawings) >= 1, "应有 draw_rect 绘制的虚线框"
    finally:
        annotated.close()


def test_burn_pdf_marks_source_regions_separately(tmp_path: Path):
    """原件侧只读取 source_page_regions，不复用回收件坐标。"""
    source_pdf = tmp_path / "source.pdf"
    target_pdf = tmp_path / "target.pdf"
    for path in (source_pdf, target_pdf):
        doc = fitz.open()
        doc.new_page(width=595, height=842)
        doc.save(path)
        doc.close()
    report = TamperReport(
        source=str(source_pdf),
        target=str(target_pdf),
        overall_risk="changed",
        change_status="changed",
        page_meta=[PageMeta(page_index=0, width_px=1654, height_px=2339, pdf_width_pt=595, pdf_height_pt=842)],
        source_page_meta=[PageMeta(page_index=0, width_px=1654, height_px=2339, pdf_width_pt=595, pdf_height_pt=842)],
        source_annotation_status="available",
        diffs=[Diff(
            alignment_id="al-source",
            status="modified",
            # target 坐标故意落在 source PDF 不存在的第 2 页；若错误复用会没有标注。
            page_regions=[PageRegion(page_index=1, bbox=[0.6, 0.6, 0.8, 0.7])],
            source_page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.3, 0.2])],
        )],
    )

    out = tmp_path / "source-annotated.pdf"
    burn_pdf(source_pdf, report, out, side="source")

    with fitz.open(out) as annotated:
        annots = list(annotated[0].annots() or [])
        assert len(annots) == 1


def test_pipeline_pdf_source_keeps_bidirectional_regions(tmp_path: Path):
    """PDF 原件经同源 OCR 取得 bbox 后，可为修改同时生成两侧标注。"""
    source_pdf = tmp_path / "source.pdf"
    target_pdf = tmp_path / "target.pdf"
    for path, text in ((source_pdf, "第一条 金额100元"), (target_pdf, "第一条 金额900元")):
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 100), text, fontname="china-s", fontsize=16)
        doc.save(path)
        doc.close()

    report = run_pipeline(
        source_pdf, target_pdf, ocr=_TextLayerOCR(), embed=MockEmbedding(),
    )

    assert report.source_annotation_status == "available"
    assert report.source_page_meta
    changed = next(diff for diff in report.diffs if diff.status == "modified")
    assert changed.page_regions
    assert changed.source_page_regions


def test_pipeline_docx_source_uses_rendered_pdf_for_real_source_regions(tmp_path: Path):
    """DOCX 仍按结构解析，比对完成后仅以派生 PDF 的真实 bbox 标注原件。"""
    word_path = tmp_path / "source.docx"
    word_path.write_bytes(_make_word([
        ("h1", "第一条 合同金额"),
        ("p", "金额为100万元。"),
    ]).getvalue())
    rendered_source = tmp_path / "source-rendered.pdf"
    target_pdf = tmp_path / "target.pdf"
    rendered_source.write_bytes(_make_pdf_from_lines([
        "第一条 合同金额",
        "金额为100万元。",
    ]).getvalue())
    target_pdf.write_bytes(_make_pdf_from_lines([
        "第一条 合同金额",
        "金额为900万元。",
    ]).getvalue())

    report = run_pipeline(
        word_path,
        target_pdf,
        source_annotation_pdf_path=rendered_source,
        ocr=_TextLayerOCR(),
        embed=MockEmbedding(),
    )

    assert report.source_annotation_status == "available"
    assert report.source_page_meta
    changed = next(diff for diff in report.diffs if diff.status == "modified")
    assert changed.source_page_regions


def test_pipeline_highlights_only_changed_paragraph_block(tmp_path: Path):
    """同一条款仅末段有差异时，两侧都不能高亮标题和未修改段落。"""
    word_path = tmp_path / "source.docx"
    word_path.write_bytes(_make_word([
        ("h1", "第一条 付款方式"),
        ("p", "第一款 付款日期为2026年1月1日。"),
        ("p", "第二款 付款金额为100万元。"),
    ]).getvalue())
    rendered_source = tmp_path / "source-rendered.pdf"
    target_pdf = tmp_path / "target.pdf"
    rendered_source.write_bytes(_make_pdf_from_lines([
        "第一条 付款方式",
        "第一款 付款日期为2026年1月1日。",
        "第二款 付款金额为100万元。",
    ]).getvalue())
    target_pdf.write_bytes(_make_pdf_from_lines([
        "第一条 付款方式",
        "第一款 付款日期为2026年1月1日。",
        "第二款 付款金额为900万元。",
    ]).getvalue())

    report = run_pipeline(
        word_path,
        target_pdf,
        source_annotation_pdf_path=rendered_source,
        ocr=_TextLayerOCR(),
        embed=MockEmbedding(),
    )

    changed = next(diff for diff in report.diffs if diff.status == "modified")
    assert len(changed.source_page_regions) == 1
    assert len(changed.page_regions) == 1
    # 第三行的真实 bbox；若错误包含标题/第一款，顶部会落在 0.12 之前。
    assert changed.source_page_regions[0].bbox[1] > 0.11
    assert changed.page_regions[0].bbox[1] > 0.11


def test_render_docx_to_pdf_runs_headless_libreoffice(monkeypatch, tmp_path: Path):
    """渲染器只产生派生 PDF，且调用固定的 LibreOffice Writer 导出参数。"""
    import subprocess

    source = tmp_path / "source.docx"
    source.write_bytes(b"not parsed by mocked libreoffice")
    output = tmp_path / "reports" / "source-rendered.pdf"

    def _fake_run(command, **kwargs):
        out_dir = Path(command[command.index("--outdir") + 1])
        (out_dir / "source.pdf").write_bytes(b"%PDF-1.4 mocked")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr("document_comparison.parsing.word_render.subprocess.run", _fake_run)
    actual = render_docx_to_pdf(source, output, executable="soffice-test", timeout_seconds=12)

    assert actual == output
    assert output.read_bytes().startswith(b"%PDF")


class _TextLayerOCR:
    """用 PDF 文本层充当 mock OCR(避开真实 LLM 调用,测试稳定可复现)。"""

    def recognize(self, pdf_path: Path, page_metas: list[PageMeta], *, on_progress=None):
        return extract_text_blocks(pdf_path)


class _MisindexedTextLayerOCR:
    """模拟第三方 OCR 把多个物理页错误标成同一页的回归用例。"""

    def recognize(self, pdf_path: Path, page_metas: list[PageMeta], *, on_progress=None):
        return [
            [block.model_copy(update={"page_index": 1}) for block in blocks]
            for blocks in extract_text_blocks(pdf_path)
        ]


class _MergedTextLayerOCR:
    """模拟原生 PDF 把多个逻辑条款合并在同一文本块的情况。"""

    def recognize(self, pdf_path: Path, page_metas: list[PageMeta], *, on_progress=None):
        return [[
            Block(
                block_id="merged",
                page_index=0,
                label="text",
                bbox=[72, 72, 520, 180],
                content=(
                    "第三条费用及支付方式 3.1 总额为100元。"
                    "3.2 支付50%预付款。"
                ),
            )
        ]]


def _make_word(parts):
    doc = Document()
    for kind, text in parts:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "p":
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf


def _make_pdf_from_lines(lines):
    # 用内置 CJK 字体 china-s,确保中文写入文本层并可被读回
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    y = 72
    for ln in lines:
        page.insert_text((72, y), ln, fontname="china-s", fontsize=11)
        y += 20
    buf = io.BytesIO()
    pdf.save(buf)
    buf.seek(0)
    return buf


def _to_files(word_buf, pdf_buf, tmp_path):
    wpath = tmp_path / "c.docx"
    ppath = tmp_path / "c.pdf"
    wpath.write_bytes(word_buf.getvalue())
    ppath.write_bytes(pdf_buf.getvalue())
    word_buf.seek(0)
    return wpath, ppath


def test_parse_source_dispatches_by_extension(tmp_path):
    """_parse_source 按后缀分发:.docx 走 parse_word(word);.pdf 走 OCR→blocks_to_raw(pdf)。"""
    from document_comparison.config import settings

    # .docx → doc_type="word"
    wpath = tmp_path / "c.docx"
    wpath.write_bytes(_make_word([("h1", "第一条 标的")]).getvalue())
    raw, doc_type = _parse_source(wpath, _TextLayerOCR(), settings, on_progress=None)
    assert doc_type == "word"
    assert any("标的" in item.text for item in raw)

    # .pdf → doc_type="pdf",经 OCR 文本层抽取得到 RawItem
    ppath = tmp_path / "c.pdf"
    ppath.write_bytes(_make_pdf_from_lines(["第一条 标的"]).getvalue())
    raw, doc_type = _parse_source(ppath, _TextLayerOCR(), settings, on_progress=None)
    assert doc_type == "pdf"
    assert any("标的" in item.text for item in raw)


def test_pipeline_accepts_pdf_source(tmp_path):
    """原始合同为 PDF 时,端到端走 OCR 解析分支并完成比对。"""
    # source 与 target 均为带文本层的 PDF,内容一致 → change_status=clean。
    pdf_buf = _make_pdf_from_lines(["第一条 合同标的", "本合同自签字之日起生效。"])
    spath = tmp_path / "source.pdf"
    tpath = tmp_path / "target.pdf"
    spath.write_bytes(pdf_buf.getvalue())
    tpath.write_bytes(pdf_buf.getvalue())
    report = run_pipeline(spath, tpath, ocr=_TextLayerOCR(), embed=MockEmbedding())
    assert isinstance(report, TamperReport)
    assert report.change_status == "clean"


def test_pipeline_identical(tmp_path):
    parts = [
        ("h1", "第一条 合同标的"),
        ("p", "甲方提供成套设备。"),
        ("h1", "第二条 合同金额"),
        ("p", "金额为100万元。"),
    ]
    word_buf = _make_word(parts)
    pdf_buf = _make_pdf_from_lines([
        "第一条 合同标的",
        "甲方提供成套设备。",
        "第二条 合同金额",
        "金额为100万元。",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    report = run_pipeline(wpath, ppath)
    assert report.overall_risk == "clean"
    assert len(report.diffs) == 0


def test_pipeline_uses_raw_span_plan_when_llm_alignment_enabled(
    tmp_path,
    monkeypatch,
):
    """开启后先联合规划 RawSpan，关闭时不得产生隐式模型调用。"""
    from document_comparison import pipeline as pipeline_module
    from document_comparison.models import (
        RawAlignmentGroup,
        RawAlignmentPlan,
        RawSpan,
    )

    preamble = "甲乙双方本着平等互利,诚实信用的原则,经过协商一致订立本合同,恪守合同规定。"
    word_buf = _make_word([
        ("p", "乙方：(供方)湖北欧朗机械有限公司"),
        ("p", preamble),
        ("p", "一 、 合同标的"),
    ])
    pdf_buf = _make_pdf_from_lines([
        "乙方：(供方)湖北欧朗机械有限公司",
        preamble,
        "一、合同标的",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    calls: list[tuple[int, int]] = []

    def planner(word_blocks, pdf_blocks):
        calls.append((len(word_blocks), len(pdf_blocks)))
        assert len(word_blocks) == len(pdf_blocks) == 3
        return RawAlignmentPlan(
            groups=[
                RawAlignmentGroup(
                    word_spans=[
                        RawSpan(block_id=f"w-{index}", start=0, end=len(word.text))
                    ],
                    pdf_spans=[
                        RawSpan(block_id=f"p-{index}", start=0, end=len(pdf.text))
                    ],
                    confidence=0.99,
                    reason="对应原始段落",
                )
                for index, (word, pdf) in enumerate(
                    zip(word_blocks, pdf_blocks), start=1
                )
            ]
        )

    monkeypatch.setattr(
        pipeline_module,
        "resolve_raw_alignment_plan",
        planner,
        raising=False,
    )
    monkeypatch.setattr(
        pipeline_module,
        "build_clauses",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("成功的 RawSpan 联合计划不应依赖 Clause 切分")
        ),
    )

    report = run_pipeline(
        wpath,
        ppath,
        ocr=_TextLayerOCR(),
        embed=MockEmbedding(),
        enable_llm_alignment=True,
    )

    assert calls == [(3, 3)]
    assert report.change_status == "clean"
    assert report.diffs == []
    assert report.unmatched_clauses == []


def test_pipeline_raw_span_plan_cannot_clear_amount_change(tmp_path, monkeypatch):
    """联合计划只决定配对；金额字符变化仍由确定性裁决报告。"""
    from document_comparison import pipeline as pipeline_module
    from document_comparison.models import (
        RawAlignmentGroup,
        RawAlignmentPlan,
        RawSpan,
    )

    word_buf = _make_word([
        ("h1", "第二条 合同金额"),
        ("p", "金额为100万元。"),
    ])
    pdf_buf = _make_pdf_from_lines([
        "第二条 合同金额",
        "金额为200万元。",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)

    def planner(word_blocks, pdf_blocks):
        return RawAlignmentPlan(
            groups=[
                RawAlignmentGroup(
                    word_spans=[
                        RawSpan(
                            block_id=f"w-{index}",
                            start=0,
                            end=len(word.text),
                        )
                    ],
                    pdf_spans=[
                        RawSpan(
                            block_id=f"p-{index}",
                            start=0,
                            end=len(pdf.text),
                        )
                    ],
                    confidence=0.99,
                    reason="同一位置段落",
                )
                for index, (word, pdf) in enumerate(
                    zip(word_blocks, pdf_blocks), start=1
                )
            ]
        )

    monkeypatch.setattr(pipeline_module, "resolve_raw_alignment_plan", planner)

    report = run_pipeline(
        wpath,
        ppath,
        ocr=_TextLayerOCR(),
        embed=MockEmbedding(),
        enable_llm_alignment=True,
        enable_risk_assessment=True,
    )

    assert report.change_status == "changed"
    assert any(diff.status == "modified" for diff in report.diffs)
    assert any(element.kind == "amount" for element in report.key_elements)


def test_pipeline_invalid_raw_plan_falls_back_to_clause_alignment(
    tmp_path,
    monkeypatch,
):
    from document_comparison import pipeline as pipeline_module

    word_buf = _make_word([("h1", "第一条 合同标的")])
    pdf_buf = _make_pdf_from_lines(["第一条 合同标的"])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    build_calls: list[str] = []
    fallback_resolvers: list[object | None] = []
    real_build = pipeline_module.build_clauses
    real_align = pipeline_module.align_clauses

    def spy_build(items, doc_type):
        build_calls.append(doc_type)
        return real_build(items, doc_type)

    def spy_align(*args, **kwargs):
        fallback_resolvers.append(kwargs.get("llm_resolver"))
        return real_align(*args, **kwargs)

    monkeypatch.setattr(
        pipeline_module,
        "resolve_raw_alignment_plan",
        lambda _word, _pdf: None,
    )
    monkeypatch.setattr(pipeline_module, "build_clauses", spy_build)
    monkeypatch.setattr(pipeline_module, "align_clauses", spy_align)

    report = run_pipeline(
        wpath,
        ppath,
        ocr=_TextLayerOCR(),
        embed=MockEmbedding(),
        enable_llm_alignment=True,
    )

    assert build_calls == ["word", "pdf"]
    assert fallback_resolvers == [None]
    assert report.change_status == "clean"


def test_pipeline_detects_amount_tamper(tmp_path):
    parts = [
        ("h1", "第二条 合同金额"),
        ("p", "金额为100万元。"),
    ]
    word_buf = _make_word(parts)
    pdf_buf = _make_pdf_from_lines([
        "第二条 合同金额",
        "金额为200万元。",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    report = run_pipeline(wpath, ppath, enable_risk_assessment=True)
    assert report.overall_risk == "high"
    # 金额变化应体现在 key_elements
    assert any(e.kind == "amount" and e.changed for e in report.key_elements)


def test_pipeline_risk_disabled_by_default(tmp_path):
    """默认 enable_risk_assessment=False:同样的金额篡改,差异仍被报告,
    但不做风险分级、不抽取 key_elements,overall 镜像 change_status(=changed,不再 low)。"""
    parts = [
        ("h1", "第二条 合同金额"),
        ("p", "金额为100万元。"),
    ]
    word_buf = _make_word(parts)
    pdf_buf = _make_pdf_from_lines([
        "第二条 合同金额",
        "金额为200万元。",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    # 不传 enable_risk_assessment(用默认 False)
    report = run_pipeline(wpath, ppath)

    # 差异仍被识别
    assert len(report.diffs) >= 1
    assert all(d.status == "modified" for d in report.diffs)
    # 风险字段被短路
    assert all(d.risk_level == "none" for d in report.diffs)
    assert all(d.risk_reasons == [] for d in report.diffs)
    # 高风险要素未抽取
    assert report.key_elements == []
    # overall 镜像 change_status:有差异 → changed(不再出现 low/medium/high)
    assert report.change_status == "changed"
    assert report.overall_risk == "changed"


def test_pipeline_header_field_alignment_no_false_positive(tmp_path):
    """首部键值块按字段名配对,不应误报 added/deleted(注入 mock OCR/embed,不依赖网络)。"""
    parts = [
        ("p", "甲方(甲方主体):XX公司"),
        ("p", "联系地址:上海市浦东新区"),
        ("p", "联系电话:13800138000"),
        ("p", "乙方(乙方主体):YY公司"),
        ("p", "联系地址:上海市黄浦区"),
        ("p", "联系电话:13900139000"),
        ("h1", "第一条 合同标的"),
        ("p", "甲方提供设备。"),
    ]
    word_buf = _make_word(parts)
    # PDF 文本层内容与 Word 完全一致(逐行)
    pdf_buf = _make_pdf_from_lines([
        "甲方(甲方主体):XX公司",
        "联系地址:上海市浦东新区",
        "联系电话:13800138000",
        "乙方(乙方主体):YY公司",
        "联系地址:上海市黄浦区",
        "联系电话:13900139000",
        "第一条 合同标的",
        "甲方提供设备。",
    ])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    report = run_pipeline(wpath, ppath, ocr=_TextLayerOCR(), embed=MockEmbedding())
    # 首部字段全部配对成功,无 added/deleted 误报
    assert len(report.unmatched_clauses) == 0, (
        f"expected no unmatched, got: {[d.status for d in report.unmatched_clauses]}"
    )
    assert report.overall_risk == "clean"


def test_pipeline_recovers_inline_pdf_boundaries_and_marks_only_real_change(tmp_path):
    """块内条款合并时仍应配对成功，并只报告真实比例变化。"""
    word_buf = _make_word([
        ("h1", "第三条 费用及支付方式"),
        ("p", "3.1 总额为100元。"),
        ("p", "3.2 支付60%预付款。"),
    ])
    # 让 PDF 侧 3.2 变为 50%，同时模拟章节、3.1、3.2 在同一原生 block。
    pdf_buf = _make_pdf_from_lines(["placeholder"])
    wpath, ppath = _to_files(word_buf, pdf_buf, tmp_path)
    progress_events: list[tuple[str, float]] = []
    report = run_pipeline(
        wpath,
        ppath,
        ocr=_MergedTextLayerOCR(),
        embed=MockEmbedding(),
        on_progress=lambda stage, progress: progress_events.append((stage, progress)),
        enable_risk_assessment=True,
    )

    assert not report.unmatched_clauses
    assert [diff.number for diff in report.diffs] == ["3.2"]
    assert any(element.kind == "ratio" for element in report.key_elements)
    assert [stage for stage, _ in progress_events if stage in {"structure", "align", "compare"}] == [
        "structure", "align", "compare",
    ]

    annotated_pdf = tmp_path / "annotated.pdf"
    burn_pdf(ppath, report, annotated_pdf)
    assert annotated_pdf.exists() and annotated_pdf.stat().st_size > 0


def _make_multipage_pdf(tmp_path: Path, n_pages: int) -> Path:
    """生成 n 页的文本层 PDF,每页写入唯一页码文本。"""
    path = tmp_path / f"{n_pages}p.pdf"
    pdf = fitz.open()
    for i in range(n_pages):
        page = pdf.new_page(width=595, height=842)
        page.insert_text((72, 100), f"PAGE-{i}", fontname="helv", fontsize=24)
    pdf.save(path)
    pdf.close()
    return path


def test_pipeline_corrects_misindexed_ocr_blocks_before_pdf_highlighting(tmp_path: Path):
    """页 1、2、4 的差异不得因 OCR 页号错误而都画到第 2 页。"""
    wpath = tmp_path / "c.docx"
    word = Document()
    for page_number in range(1, 5):
        word.add_heading(f"第{page_number}条 原始内容", level=1)
        if page_number < 4:
            word.add_page_break()
    word.save(wpath)

    ppath = tmp_path / "4p-changed.pdf"
    pdf = fitz.open()
    for page_number in range(1, 5):
        page = pdf.new_page(width=595, height=842)
        page.insert_text(
            (72, 100),
            f"第{page_number}条 回收件已修改内容",
            fontname="china-s",
            fontsize=16,
        )
    pdf.save(ppath)
    pdf.close()

    with llm_call_collector() as audit_records:
        report = run_pipeline(
            wpath, ppath, ocr=_MisindexedTextLayerOCR(), embed=MockEmbedding(),
        )
    ocr_result = next(record for record in audit_records if record.kind == "ocr-result")
    assert isinstance(ocr_result.response, dict)
    assert {
        block["block_page_index"] for block in ocr_result.response["blocks"]
    } == {0, 1, 2, 3}
    highlighted_pages = {
        region.page_index
        for diff in [*report.diffs, *report.unmatched_clauses]
        for region in diff.page_regions
    }
    assert {0, 1, 3}.issubset(highlighted_pages)

    annotated = tmp_path / "annotated.pdf"
    burn_pdf(ppath, report, annotated)
    with fitz.open(annotated) as annotated_doc:
        assert [
            len(list(annotated_doc[page_index].annots() or []))
            for page_index in (0, 1, 3)
        ] == [1, 1, 1]


def test_pipeline_truncates_trailing_drawings_after_estimated_contract_pages(
    tmp_path: Path,
):
    """自动保留三页合同并截掉回收 PDF 后面的两页图纸。

    DOCX 页数来自保存页数/分页标记；截取后全部合同页仍必须进入 OCR 和高亮。
    """
    wpath = tmp_path / "c.docx"
    word = Document()
    word.add_heading("第一条 合同标的", level=1)
    word.add_paragraph("第一页合同正文。")
    word.add_page_break()
    word.add_paragraph("第二页合同正文。")
    word.add_page_break()
    word.add_paragraph("第三页合同正文。")
    word.save(wpath)
    ppath = _make_multipage_pdf(tmp_path, 5)

    report = run_pipeline(
        wpath, ppath, ocr=_TextLayerOCR(), embed=MockEmbedding(),
        truncate_to_original_pages=True,
        truncated_pdf_output_path=tmp_path / "compared.pdf",
    )
    assert len(report.page_meta) == 3
    assert [meta.page_index for meta in report.page_meta] == [0, 1, 2]
    assert report.truncation is not None
    assert report.truncation.original_pdf_page_count == 5
    assert report.truncation.truncated_pdf_page_count == 3
    assert report.truncation.original_doc_page_count == 3
    assert report.truncation.doc_page_count_source == "estimated"
    assert report.target == str(tmp_path / "compared.pdf")
    highlighted_pages = {
        region.page_index
        for diff in [*report.diffs, *report.unmatched_clauses]
        for region in diff.page_regions
    }
    assert highlighted_pages == {0, 1, 2}

    # 下游 burn_pdf 必须保留全部三页合同，不带后面的图纸。
    annotated = tmp_path / "annotated.pdf"
    burn_pdf(report.target, report, annotated)
    assert annotated.exists() and annotated.stat().st_size > 0
    assert count_pages(report.target) == 3
    assert count_pages(annotated) == 3
    with fitz.open(annotated) as annotated_doc:
        assert [
            len(list(page.annots() or []))
            for page in annotated_doc
        ] == [1, 1, 1]


def test_pipeline_truncate_respects_explicit_page_count(tmp_path: Path):
    """显式 original_page_count 覆盖 docx 估算:2 页 PDF 中保留 2 页中的前 1 页。"""
    word_buf = _make_word([
        ("h1", "第一条 合同标的"),
        ("p", "甲方提供设备。"),
    ])
    wpath = tmp_path / "c.docx"
    wpath.write_bytes(word_buf.getvalue())
    # 回收 PDF 2 页,显式指定原始合同 1 页
    ppath = _make_multipage_pdf(tmp_path, 2)

    report = run_pipeline(
        wpath, ppath, ocr=_TextLayerOCR(), embed=MockEmbedding(),
        truncate_to_original_pages=True, original_page_count=1,
    )
    assert len(report.page_meta) == 1
    # 显式传入页数 -> source=explicit(审计可追溯页数来源)
    assert report.truncation is not None
    assert report.truncation.original_pdf_page_count == 2
    assert report.truncation.truncated_pdf_page_count == 1
    assert report.truncation.doc_page_count_source == "explicit"


def test_pipeline_no_truncation_when_pdf_within_original_pages(tmp_path: Path):
    """回收 PDF 页数 <= 原始页数时不截取:页数不变。"""
    word_buf = _make_word([
        ("h1", "第一条 合同标的"),
        ("p", "甲方提供设备。"),
    ])
    wpath = tmp_path / "c.docx"
    wpath.write_bytes(word_buf.getvalue())
    # 估算 1 页,PDF 也 1 页 -> 不截取
    ppath = _make_multipage_pdf(tmp_path, 1)

    report = run_pipeline(
        wpath, ppath, ocr=_TextLayerOCR(), embed=MockEmbedding(),
        truncate_to_original_pages=True,
    )
    assert len(report.page_meta) == 1
    # 未发生截断 -> 无留痕记录
    assert report.truncation is None


def test_pipeline_truncate_off_by_default(tmp_path: Path):
    """默认不开启截取:回收 PDF 多页保持不变。"""
    word_buf = _make_word([("h1", "第一条 合同标的"), ("p", "甲方提供设备。")])
    wpath = tmp_path / "c.docx"
    wpath.write_bytes(word_buf.getvalue())
    ppath = _make_multipage_pdf(tmp_path, 3)

    report = run_pipeline(wpath, ppath, ocr=_TextLayerOCR(), embed=MockEmbedding())
    assert len(report.page_meta) == 3
    # 未开启截取 -> 无留痕记录
    assert report.truncation is None
