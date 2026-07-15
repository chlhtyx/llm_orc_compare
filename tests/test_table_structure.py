"""结构化表格提取测试(单元格内换行 + 水平合并)。

覆盖 _table_rows 的两类版式修复:
1. 单元格内换行(表头「数\\n量」)→ 拍平为「数 量」,不拆成两行
2. 水平合并(gridSpan)→ 按 _tc 去重,合计行不重复展平

构造方式:用 python-docx 在内存生成 docx,临时落盘后调 parse_word,
不依赖外部样本文件,保证测试可移植。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

from document_comparison.parsing.word import _table_rows, parse_word


def _make_table_docx(path: Path) -> None:
    """构造一张带换行表头 + 水平合并合计行的表格。

    结构:
      表头: 序号 | 名称 | 数\n量(单元格内换行) | 单价
      数据行1: 1 | A | 100 | 10.00
      合计行: 「合计」跨前2列 | 「人民币1000元」跨后2列(水平合并)
    """
    doc = Document()
    t = doc.add_table(rows=3, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    t.cell(0, 0).text = "序号"
    t.cell(0, 1).text = "名称"
    # 单元格内换行:这是采购合同里表头「数\\n量」的真实样式
    t.cell(0, 2).paragraphs[0].add_run("数")
    t.cell(0, 2).paragraphs[0].add_run("\n")
    t.cell(0, 2).paragraphs[0].add_run("量")
    t.cell(0, 3).text = "单价"

    # 数据行
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "A"
    t.cell(1, 2).text = "100"
    t.cell(1, 3).text = "10.00"

    # 合计行:用 gridSpan 做水平合并
    # python-docx 通过合并 cell 实现水平合并
    t.cell(2, 0).text = "合计"
    t.cell(2, 0).merge(t.cell(2, 1))
    t.cell(2, 2).text = "人民币1000元"
    t.cell(2, 2).merge(t.cell(2, 3))

    doc.save(str(path))


def test_cell_internal_newline_flattened(tmp_path):
    """单元格内换行应拍平为空格,表头不被错拆成两行。"""
    docx_path = tmp_path / "t.docx"
    _make_table_docx(docx_path)

    items = parse_word(str(docx_path))
    table_items = [it for it in items if it.kind == "table"]
    assert len(table_items) == 1
    tbl = table_items[0].table
    assert tbl is not None

    # 关键断言:表头 4 列,「数 量」作为单列存在(换行→空格,未拆行)
    assert len(tbl.headers) == 4
    assert tbl.headers[2] == "数 量"

    # 纯文本里「数 量」也应在同一行(未被 normalize 拆成两行)
    first_line = table_items[0].text.splitlines()[0]
    assert "数 量" in first_line


def test_horizontal_merge_dedup(tmp_path):
    """水平合并(gridSpan)的合计行应按 _tc 去重,不重复展平。"""
    docx_path = tmp_path / "t.docx"
    _make_table_docx(docx_path)

    items = parse_word(str(docx_path))
    tbl = [it for it in items if it.kind == "table"][0].table
    assert tbl is not None

    # 合计行应是 2 个单元格(「合计」+「人民币1000元」),而非 4 个带重复
    # rows[0] 是数据行,rows[1] 是合计行
    total_row = tbl.rows[1]
    assert len(total_row) == 2
    assert total_row[0] == "合计"
    assert total_row[1] == "人民币1000元"
    # 无重复值
    assert len(set(total_row)) == len(total_row)


def test_header_data_row_column_alignment(tmp_path):
    """表头列数应与数据行列数一致(列对齐的根本保证)。"""
    docx_path = tmp_path / "t.docx"
    _make_table_docx(docx_path)

    items = parse_word(str(docx_path))
    tbl = [it for it in items if it.kind == "table"][0].table
    assert tbl is not None

    header_cols = len(tbl.headers)
    # 至少存在一个与表头列数一致的数据行(非合并行)
    aligned_rows = [r for r in tbl.rows if len(r) == header_cols]
    assert aligned_rows, f"无任何数据行与表头列数({header_cols})一致"


def test_table_rows_unit():
    """直接单测 _table_rows:换行拍平 + 合并去重的底层保证。"""
    doc = Document()
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "序号"
    t.cell(0, 1).paragraphs[0].add_run("名")
    t.cell(0, 1).paragraphs[0].add_run("\n")
    t.cell(0, 1).paragraphs[0].add_run("称")
    t.cell(0, 2).text = "备注"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "X"
    t.cell(1, 2).text = ""

    rows = _table_rows(t)
    assert rows[0] == ["序号", "名 称", "备注"]
    assert rows[1] == ["1", "X", ""]

    # 合并去重:python-docx 的 merge() 会把被合并格的文本拼到首格(row.cells
    # 对合并区返回同一 _tc)。关键是验证去重生效——合并后该行应为 2 格而非 3 格,
    # 即 gridSpan 的后续物理格已被跳过。
    t.cell(1, 0).merge(t.cell(1, 1))
    rows = _table_rows(t)
    assert len(rows[1]) == 2
