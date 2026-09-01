"""estimate_page_count:基于 OOXML 分页标记估算 docx 页数。"""
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document  # type: ignore[import-untyped]
from docx.oxml import OxmlElement  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]

from document_comparison.parsing.word import estimate_page_count

_APP_PROPERTIES_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
)


def _save(tmp_path: Path, name: str = "d.docx") -> Path:
    path = tmp_path / name
    Document().save(path)
    return path


def _add_manual_page_break(doc: Document) -> None:
    """在末尾段落插入一个 <w:br w:type="page"/>。"""
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def _add_last_rendered_page_break(doc: Document) -> None:
    """在末尾段落插入一个 <w:lastRenderedPageBreak/>(Word 渲染分页提示)。"""
    p = doc.add_paragraph()
    run = p.add_run()
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))


def _set_saved_page_count(path: Path, page_count: int) -> None:
    with ZipFile(path) as archive:
        members = {
            info.filename: archive.read(info.filename)
            for info in archive.infolist()
        }
    root = ElementTree.fromstring(members["docProps/app.xml"])
    pages = root.find(f"{{{_APP_PROPERTIES_NS}}}Pages")
    assert pages is not None
    pages.text = str(page_count)
    members["docProps/app.xml"] = ElementTree.tostring(
        root, encoding="utf-8", xml_declaration=True
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def test_no_evidence_returns_zero(tmp_path: Path):
    """无有效保存页数(Pages=0)且无分页标记时返回 0(页数未知,调用方必须跳过截取)。"""
    path = _save(tmp_path)
    _set_saved_page_count(path, 0)
    assert estimate_page_count(path) == 0


def test_saved_word_page_count_is_used_without_break_markers(tmp_path: Path):
    path = _save(tmp_path)
    _set_saved_page_count(path, 4)

    assert estimate_page_count(path) == 4


def test_manual_page_breaks(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("第一页")
    _add_manual_page_break(doc)
    doc.add_paragraph("第二页")
    _add_manual_page_break(doc)
    doc.add_paragraph("第三页")
    path = tmp_path / "d.docx"
    doc.save(path)
    # 2 个手动分页符 -> 3 页
    assert estimate_page_count(path) == 3


def test_last_rendered_page_breaks(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("内容")
    _add_last_rendered_page_break(doc)
    doc.add_paragraph("更多内容")
    _add_last_rendered_page_break(doc)
    path = tmp_path / "d.docx"
    doc.save(path)
    # 2 个渲染分页提示 -> 3 页
    assert estimate_page_count(path) == 3


def test_mixed_breaks(tmp_path: Path):
    doc = Document()
    _add_manual_page_break(doc)
    _add_last_rendered_page_break(doc)
    path = tmp_path / "d.docx"
    doc.save(path)
    # 2 个分页标记(1 手动 + 1 渲染)-> 3 页
    assert estimate_page_count(path) == 3


def test_page_count_uses_larger_of_saved_metadata_and_break_markers(
    tmp_path: Path,
):
    doc = Document()
    _add_manual_page_break(doc)
    _add_manual_page_break(doc)
    path = tmp_path / "d.docx"
    doc.save(path)
    _set_saved_page_count(path, 2)

    assert estimate_page_count(path) == 3
