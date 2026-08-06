"""docx 解析日志测试:验证 parse_word 输出的解析摘要日志。

覆盖新增的 `_log_parse_summary`:
- INFO「word parsed ...」记录各类块计数、标题层级、表格规模、字符数、耗时
- 当文档含表格但全部被空跳过时,输出 WARNING

构造方式:用 python-docx 在内存生成 docx,临时落盘后调 parse_word,
与 test_table_structure.py 同一套路,不依赖外部样本文件。
"""
from __future__ import annotations

import logging
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from document_comparison.parsing.word import parse_word
from document_comparison.structure.clause import build_clauses


def _make_rich_docx(path: Path) -> None:
    """构造含标题 + 段落 + 表格的混合文档。

    结构:
      - 标题(Heading 1):第一条 合同主体
      - 段落:甲方:... 乙方:... (键值块,但解析层只产 paragraph)
      - 标题(Heading 2):1.1 价款
      - 表格:2 行 × 2 列
      - 空段落(验证 skipped_empty 计数)
    """
    doc = Document()
    doc.add_heading("第一条 合同主体", level=1)
    doc.add_paragraph("甲方:买方公司")
    doc.add_paragraph("乙方:卖方公司")
    doc.add_heading("1.1 价款", level=2)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "金额"
    t.cell(1, 0).text = "货款"
    t.cell(1, 1).text = "1000元"
    doc.add_paragraph("")  # 空段落,应被跳过并计入 skipped_empty
    doc.save(str(path))


def test_parse_summary_info_logged(tmp_path, caplog):
    """parse_word 应输出含各类块计数与规模的 INFO 摘要日志。"""
    docx_path = tmp_path / "rich.docx"
    _make_rich_docx(docx_path)

    with caplog.at_level(logging.INFO, logger="document_comparison.parsing.word"):
        items = parse_word(str(docx_path))

    summary = next(
        (r for r in caplog.records if r.message.startswith("word parsed")), None
    )
    assert summary is not None, "应输出 word parsed 摘要日志"
    msg = summary.message

    # 文件名(不含路径,避免日志泄露绝对路径)
    assert f"path={docx_path.name}" in msg
    # 各类块计数:2 个标题 + 2 段落 + 1 表格 = 5 项
    assert "items=5" in msg
    assert "headings=2" in msg
    # 段落数取自 doc.paragraphs(python-docx 中标题也计入段落)
    # 2 标题 + 2 键值段 + 1 空段 = 5
    assert "paragraphs=5" in msg
    # 表格规模:rows×headers,rows 仅数据行(表头单独存 headers)
    # 2×2 物理表 = 1 数据行 × 2 列
    assert "table_dims=1x2" in msg
    # 跳过的空段落数
    assert "skipped_empty=1" in msg
    # 耗时字段存在
    assert "elapsed=" in msg
    # 标题层级分布:level 1 与 level 2 各一
    assert "heading_levels=" in msg

    # 返回值本身不受日志影响
    assert len(items) == 5
    assert sum(1 for it in items if it.kind == "table") == 1


def test_empty_table_dropped_warning(tmp_path, caplog):
    """文档含表格但全部为空时,应输出 WARNING 提示。"""
    doc = Document()
    # 全空表格:_table_to_text 返回空串,parse_word 会跳过它
    t = doc.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = ""
    # 补一个正常段落,使 items 非全空(单纯验证 warning 触发条件)
    doc.add_paragraph("正文")
    docx_path = tmp_path / "empty_table.docx"
    doc.save(str(docx_path))

    with caplog.at_level(logging.WARNING, logger="document_comparison.parsing.word"):
        items = parse_word(str(docx_path))

    warnings = [r for r in caplog.records if r.levelno == logging.WARNING]
    assert any("解析后 0 张入库" in r.message for r in warnings), \
        "全空表格应触发 WARNING"

    # 解析结果:只有段落,表格被跳过
    assert len(items) == 1
    assert items[0].kind == "paragraph"


def test_word_page_break_is_preserved_as_deleted_location_hint(tmp_path):
    """显式分页后的 Word 条款应携带页序，供 deleted 跨页占位定位使用。"""
    doc = Document()
    doc.add_paragraph("第十条 其他约定")
    doc.add_page_break()
    doc.add_paragraph("10.4 测试条款1232456767")
    docx_path = tmp_path / "paged.docx"
    doc.save(str(docx_path))

    items = parse_word(str(docx_path))
    clauses = build_clauses(items, "word")

    assert [item.page_index for item in items] == [0, 1]
    assert clauses[1].number == "10.4"
    assert clauses[1].source_page_index == 1


# —— 修订痕迹(track changes)处理:带修订的 docx 应按「接受修订」视图解析 ——


def _inject_revisions_into_docx(path: Path) -> Path:
    """在已保存的 docx 内注入 OOXML 修订标记,返回新文件路径。

    python-docx 无添加修订的 API,这里直接改写 ``word/document.xml``(zipfile 重打包)。
    注入内容:
    - 段落「合同金额为100万元」→「合同金额为」+ del(100万元) + ins(200万元)
      (接受视图应为「合同金额为200万元」)
    - 表格单元格「100」→ del(100) + ins(200)(接受视图应为「200」)
    """
    src = path
    with zipfile.ZipFile(src) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    # 段落修订:删除「100万元」、插入「200万元」
    doc_xml = doc_xml.replace(
        "<w:t>合同金额为100万元</w:t>",
        "<w:r><w:t>合同金额为</w:t></w:r>"
        '<w:del w:id="100" w:author="alice" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:delText>100万元</w:delText></w:r></w:del>"
        '<w:ins w:id="101" w:author="alice" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:t>200万元</w:t></w:r></w:ins>",
        1,
    )
    # 表格单元格修订:删除「100」、插入「200」
    doc_xml = doc_xml.replace(
        "<w:t>100</w:t>",
        '<w:del w:id="200" w:author="alice" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:delText>100</w:delText></w:r></w:del>"
        '<w:ins w:id="201" w:author="alice" w:date="2024-01-01T00:00:00Z">'
        "<w:r><w:t>200</w:t></w:r></w:ins>",
        1,
    )

    out = src.with_suffix(".revised.docx")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = doc_xml.encode("utf-8")
            zout.writestr(item, data)
    return out


def _make_docx_with_revisions(path: Path) -> Path:
    """构造一份带未接受修订痕迹的 docx,返回落盘路径(含修订版)。

    结构:
      - 段落「合同金额为100万元」(会被改写为含 del/ins 的修订段)
      - 表格 1×2:表头「项目|金额」,数据行「货款|100」(单元格 100 被改写为修订)
    """
    doc = Document()
    doc.add_paragraph("合同金额为100万元")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "金额"
    t.cell(1, 0).text = "货款"
    t.cell(1, 1).text = "100"
    base = path.with_suffix(".docx")
    doc.save(str(base))
    return _inject_revisions_into_docx(base)


def test_parse_word_accepts_revisions_in_paragraph(tmp_path):
    """回归 #1:带修订的段落应产出「接受修订」视图,而非空串。

    历史 bug:python-docx ``paragraph.text`` 只取段落直接 ``<w:r>`` 子元素,
    被 ``<w:ins>``/``<w:del>`` 包裹的 run 不可见 → 带修订段落被整段丢弃成 ``''``,
    Word 侧基准文本凭空消失,与 PDF 比对产生伪差异/漏检。
    修复后:删除痕迹(``<w:delText>``)丢弃,插入痕迹(``<w:ins>`` 内 ``<w:t>``)并入。
    """
    revised = _make_docx_with_revisions(tmp_path / "revised")
    items = parse_word(str(revised))

    para_texts = [it.text for it in items if it.kind == "paragraph"]
    # 接受视图:「合同金额为」+「200万元」(插入),「100万元」(删除)被丢弃
    assert "合同金额为200万元" in para_texts
    # 删除痕迹文本不应出现
    assert all("100万元" not in t for t in para_texts)
    # 关键:不再因修订丢成空串
    assert para_texts == ["合同金额为200万元"]


def test_parse_word_accepts_revisions_in_table_cell(tmp_path):
    """回归 #2:含修订的表格单元格应产出「接受修订」视图。

    单元格「100」被改写为 del(100)+ins(200);接受视图应为「200」。
    """
    revised = _make_docx_with_revisions(tmp_path / "revised")
    items = parse_word(str(revised))

    table_item = next(it for it in items if it.kind == "table")
    assert table_item.table is not None
    # 数据行「货款 | 200」(100 被删除痕迹丢弃,200 被插入痕迹并入)
    flat = " | ".join(table_item.table.headers) + " | " + " | ".join(
        c for row in table_item.table.rows for c in row
    )
    assert "200" in flat
    assert "100" not in flat  # 删除痕迹文本不应残留


def test_parse_word_no_revision_unchanged(tmp_path):
    """回归:无修订的普通 docx 文本与新逻辑一致(确保 _accepted_text 不破坏既有行为)。"""
    docx_path = tmp_path / "plain.docx"
    _make_rich_docx(docx_path)
    items = parse_word(str(docx_path))

    # 与既有测试同构:2 标题 + 2 段落 + 1 表格
    assert len(items) == 5
    para_texts = [it.text for it in items if it.kind == "paragraph"]
    assert "甲方:买方公司" in para_texts
    assert "乙方:卖方公司" in para_texts


def test_parse_summary_logs_revision_count(tmp_path, caplog):
    """带修订的 docx 解析摘要日志应含 revisions=N(N>0),便于运维识别。"""
    revised = _make_docx_with_revisions(tmp_path / "revised")
    with caplog.at_level(logging.INFO, logger="document_comparison.parsing.word"):
        parse_word(str(revised))

    summary = next(
        (r for r in caplog.records if r.message.startswith("word parsed")), None
    )
    assert summary is not None
    # 注入了 4 处修订标记(2 段 del/ins + 2 单元格 del/ins)
    assert "revisions=4" in summary.message

