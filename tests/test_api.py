"""API 端点冒烟测试(使用 FastAPI TestClient)。"""
import pytest
from fastapi.testclient import TestClient

from document_comparison.api.app import create_app
from document_comparison.config import settings


@pytest.fixture
def client(db_isolated):
    """所有 API 测试在隔离的测试库中跑(避免污染开发库 doc_compare)。

    db_isolated fixture 已 monkeypatch 了 settings.database_url 与引擎工厂,
    所以 create_app() 内的 db_pkg.init_engine() 会连到测试库。
    """
    return TestClient(create_app())


@pytest.fixture(autouse=True)
def _reset_task_manager():
    """每个测试前后清空 task_manager,避免后台任务跨测试残留导致并发上限误触发。"""
    from document_comparison.api.app import task_manager
    task_manager._tasks.clear()
    yield
    task_manager._tasks.clear()


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_scan_protection_rejects_sensitive_path_and_adds_security_headers(client):
    """敏感文件探测不进入路由，正常响应也带基础浏览器防护头。"""
    blocked = client.get("/.env")
    assert blocked.status_code == 404
    assert blocked.json() == {"code": 404, "message": "not found"}
    absolute_path = client.request("GET", "//etc/passwd")
    assert absolute_path.status_code == 404
    assert "root:x:0:0" not in absolute_path.text
    assert blocked.headers["x-content-type-options"] == "nosniff"
    assert blocked.headers["x-frame-options"] == "DENY"

    healthy = client.get("/health")
    assert healthy.headers["referrer-policy"] == "same-origin"
    assert "camera=()" in healthy.headers["permissions-policy"]


def test_version_endpoint(client):
    """GET /api/v1/version 始终返回一个非空版本字符串(真相源:settings.version)。"""
    r = client.get("/api/v1/version")
    assert r.status_code == 200
    body = r.json()
    assert isinstance(body["version"], str)
    assert body["version"]


def test_version_endpoint_follows_dc_version(client, monkeypatch):
    """DC_VERSION 是真正的单一真相源:改 settings.version 端点立即跟随。

    镜像 tag(${DC_VERSION})与前端显示应同源,改一处即同步。
    """
    monkeypatch.setattr(settings, "version", "9.9.9")
    r = client.get("/api/v1/version")
    assert r.status_code == 200
    assert r.json()["version"] == "9.9.9"


def test_compare_rejects_non_docx_non_pdf_source(client):
    import io

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("bad.txt", io.BytesIO(b"text"), "text/plain"),
            "target": ("ok.pdf", io.BytesIO(b"%PDF"), "application/pdf"),
        },
    )
    assert r.status_code == 400
    msg = r.json()["message"]
    assert "docx" in msg and "pdf" in msg


def test_compare_accepts_pdf_source(client):
    """原始合同允许 PDF(.pdf)上传(与 .docx 并列)。"""
    import io

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.pdf", io.BytesIO(_make_pdf_bytes(1)), "application/pdf"),
            "target": ("scan.pdf", io.BytesIO(_make_pdf_bytes(1)), "application/pdf"),
        },
    )
    assert r.status_code == 200
    assert "task_id" in r.json()


def test_compare_accepts_valid_files(client):
    import io
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 72), "第一条 测试条款", fontsize=12)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", doc_buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", pdf_buf, "application/pdf"),
        },
    )
    assert r.status_code == 200
    assert "task_id" in r.json()


def test_compare_accepts_llm_alignment_and_risk_options(client, monkeypatch):
    """LLM 对齐和风险说明是相互独立的提交选项，并应完整透传。"""
    import io
    from docx import Document  # type: ignore[import-untyped]

    from document_comparison.api.app import task_manager

    # 捕获 run 调用参数,避免真实 OCR
    captured: dict = {}
    async def _spy_run(*args, **kwargs):
        captured.update(kwargs)
    monkeypatch.setattr(task_manager, "run", _spy_run)

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    pdf.new_page(width=595, height=842)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", doc_buf,
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", pdf_buf, "application/pdf"),
        },
        data={
            "options": (
                '{"enable_llm_alignment": true, '
                '"enable_risk_assessment": true, "enable_llm_judge": true}'
            )
        },
    )
    assert r.status_code == 200, r.json()
    assert captured.get("enable_risk_assessment") is True
    assert captured.get("enable_llm_judge") is True
    assert captured.get("enable_llm_alignment") is True


def test_compare_rejects_bad_options(client):
    """options 非法 JSON / 类型应返回 400。"""
    import io

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("bad.docx", io.BytesIO(b"x"), "application/octet-stream"),
            "target": ("ok.pdf", io.BytesIO(b"%PDF"), "application/pdf"),
        },
        data={"options": "not-json"},
    )
    assert r.status_code == 400


def test_compare_accepts_truncate_options(client, monkeypatch):
    """options 携带 truncate_to_original_pages / original_page_count 应透传到 run。"""
    import io
    from docx import Document  # type: ignore[import-untyped]

    from document_comparison.api.app import task_manager

    captured: dict = {}

    async def _spy_run(*args, **kwargs):
        captured.update(kwargs)

    monkeypatch.setattr(task_manager, "run", _spy_run)

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    pdf.new_page(width=595, height=842)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", doc_buf,
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", pdf_buf, "application/pdf"),
        },
        data={
            "options": (
                '{"truncate_to_original_pages": true, '
                '"original_page_count": 5}'
            ),
        },
    )
    assert r.status_code == 200, r.json()
    assert captured.get("truncate_to_original_pages") is True
    assert captured.get("original_page_count") == 5


def test_preview_and_annotated_report_use_truncated_task_pdf(client, monkeypatch, tmp_path):
    """报告页的预览与标注下载不应重新返回原始全页回收件。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    from document_comparison.api.app import task_manager
    from document_comparison.models import TamperReport, TruncationRecord
    from document_comparison.storage import compared_pdf_path

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare")
    task = task_manager.get(task_id)
    assert task is not None

    original = settings.uploads_dir / f"{task_id}-target.pdf"
    compared = compared_pdf_path(task_id)
    for path, page_count in ((original, 3), (compared, 1)):
        pdf = fitz.open()
        for _ in range(page_count):
            pdf.new_page(width=300, height=400)
        pdf.save(path)
        pdf.close()

    task.report = TamperReport(
        source="source.docx",
        target=str(compared),
        truncation=TruncationRecord(
            original_pdf_page_count=3,
            truncated_pdf_page_count=1,
            original_doc_page_count=1,
            doc_page_count_source="explicit",
        ),
    )

    preview = client.get(f"/api/v1/compare/{task_id}/source")
    annotated = client.get(f"/api/v1/compare/{task_id}/report?format=pdf")
    assert preview.status_code == 200
    assert annotated.status_code == 200
    with fitz.open(stream=preview.content, filetype="pdf") as doc:
        assert len(doc) == 1
    with fitz.open(stream=annotated.content, filetype="pdf") as doc:
        assert len(doc) == 1


def test_html_report_download_embeds_highlight_pages_and_clause_navigation(client, monkeypatch, tmp_path):
    """Web HTML 导出复用已验证坐标，表格行可跳转到对应高亮页。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    from document_comparison.api.app import task_manager
    from document_comparison.models import Diff, PageRegion, TamperReport

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare", document_no="WEB-HTML-001")
    task = task_manager.get(task_id)
    assert task is not None
    for role in ("source", "target"):
        path = settings.uploads_dir / f"{task_id}-{role}.pdf"
        pdf = fitz.open()
        pdf.new_page(width=300, height=400)
        pdf.save(path)
        pdf.close()
    task.report = TamperReport(
        source="source.pdf",
        target="target.pdf",
        change_status="changed",
        source_annotation_status="available",
        diffs=[Diff(
            alignment_id="html-nav",
            status="modified",
            page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
            source_page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
        )],
    )

    response = client.get(f"/api/v1/compare/{task_id}/report?format=html")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "attachment;" in response.headers["content-disposition"]
    assert 'data-target-page="1"' in response.text
    assert 'data-source-page="1"' in response.text
    assert 'id="target-page-1"' in response.text
    assert 'id="source-page-1"' in response.text
    assert "row.dataset.targetPage" in response.text


def test_pdf_report_download_generates_then_reuses(client, monkeypatch, tmp_path):
    """report.pdf 首次下载按需生成(与 HTML 导出同款管线),再次下载复用盘上产物。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    from document_comparison.api.app import task_manager
    from document_comparison.external_api import external_pdf_report_path
    from document_comparison.models import Diff, PageRegion, TamperReport

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare", document_no="WEB-PDF-001")
    task = task_manager.get(task_id)
    assert task is not None
    for role in ("source", "target"):
        path = settings.uploads_dir / f"{task_id}-{role}.pdf"
        pdf = fitz.open()
        pdf.new_page(width=300, height=400)
        pdf.save(path)
        pdf.close()
    task.report = TamperReport(
        source="source.pdf",
        target="target.pdf",
        change_status="changed",
        source_annotation_status="available",
        diffs=[Diff(
            alignment_id="pdf-dl",
            status="modified",
            page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
            source_page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
        )],
    )

    response = client.get(f"/api/v1/compare/{task_id}/report.pdf")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    disposition = response.headers["content-disposition"]
    assert "attachment;" in disposition
    assert "filename*=UTF-8''" in disposition
    assert disposition.endswith(".pdf")
    assert response.content.startswith(b"%PDF")
    with fitz.open(stream=response.content, filetype="pdf") as doc:
        assert len(doc) >= 1
    assert external_pdf_report_path(task_id).is_file()

    again = client.get(f"/api/v1/compare/{task_id}/report.pdf")
    assert again.status_code == 200
    assert again.content == response.content


def test_pdf_report_download_404_without_report(client, monkeypatch, tmp_path):
    """任务无报告(未完成/未知)时下载返回 404。"""
    from document_comparison.api.app import task_manager

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare")
    r = client.get(f"/api/v1/compare/{task_id}/report.pdf")
    assert r.status_code == 404
    unknown = client.get("/api/v1/compare/does-not-exist/report.pdf")
    assert unknown.status_code == 404


def test_original_pdf_preview_and_source_annotated_report(client, monkeypatch, tmp_path):
    """原件为 PDF 时，内部接口能独立预览和下载原件侧标注。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    from document_comparison.api.app import task_manager
    from document_comparison.models import Diff, PageMeta, PageRegion, TamperReport

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare")
    task = task_manager.get(task_id)
    assert task is not None
    source = settings.uploads_dir / f"{task_id}-source.pdf"
    target = settings.uploads_dir / f"{task_id}-target.pdf"
    for path in (source, target):
        pdf = fitz.open()
        pdf.new_page(width=300, height=400)
        pdf.save(path)
        pdf.close()
    page_meta = PageMeta(
        page_index=0, width_px=300, height_px=400, pdf_width_pt=300, pdf_height_pt=400,
    )
    task.report = TamperReport(
        source=str(source), target=str(target), change_status="changed",
        page_meta=[page_meta], source_page_meta=[page_meta],
        source_annotation_status="available",
        diffs=[Diff(
            alignment_id="source-diff", status="modified",
            source_page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
        )],
    )

    preview = client.get(f"/api/v1/compare/{task_id}/original-pdf")
    annotated = client.get(f"/api/v1/compare/{task_id}/report?format=pdf&side=source")
    assert preview.status_code == 200
    assert annotated.status_code == 200
    with fitz.open(stream=annotated.content, filetype="pdf") as doc:
        assert len(doc) == 1
        assert len(list(doc[0].annots() or [])) == 1


def test_docx_rendered_pdf_preview_and_source_annotated_report(client, monkeypatch, tmp_path):
    """DOCX 原件预览/标注使用任务生成的派生 PDF，而不是改写上传 DOCX。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    from document_comparison.api.app import task_manager
    from document_comparison.models import Diff, PageMeta, PageRegion, TamperReport

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare")
    task = task_manager.get(task_id)
    assert task is not None
    source = settings.uploads_dir / f"{task_id}-source.docx"
    target = settings.uploads_dir / f"{task_id}-target.pdf"
    source.write_bytes(b"uploaded docx remains unchanged")
    for path in (target, settings.reports_dir / f"{task_id}_source_rendered.pdf"):
        pdf = fitz.open()
        pdf.new_page(width=300, height=400)
        pdf.save(path)
        pdf.close()
    page_meta = PageMeta(
        page_index=0, width_px=300, height_px=400, pdf_width_pt=300, pdf_height_pt=400,
    )
    task.report = TamperReport(
        source=str(source), target=str(target), change_status="changed",
        page_meta=[page_meta], source_page_meta=[page_meta],
        source_annotation_status="available",
        diffs=[Diff(
            alignment_id="docx-source-diff", status="modified",
            source_page_regions=[PageRegion(page_index=0, bbox=[0.1, 0.1, 0.5, 0.2])],
        )],
    )

    preview = client.get(f"/api/v1/compare/{task_id}/original-pdf")
    annotated = client.get(f"/api/v1/compare/{task_id}/report?format=pdf&side=source")
    assert preview.status_code == 200
    assert annotated.status_code == 200
    assert source.read_bytes() == b"uploaded docx remains unchanged"


def test_history_source_download_returns_uploaded_original_file(client, monkeypatch, tmp_path):
    """比对记录下载的是原上传 DOCX，不是原件侧预览用的派生 PDF。"""
    from document_comparison.api.app import task_manager

    monkeypatch.setattr(settings, "storage_dir", tmp_path / "storage")
    settings.ensure_dirs()
    task_id = task_manager.create("compare", source_name="采购合同.docx")
    source = settings.uploads_dir / f"{task_id}-source.docx"
    source.write_bytes(b"original-docx-bytes")

    response = client.get(f"/api/v1/tasks/{task_id}/source/download")

    assert response.status_code == 200
    assert response.content == b"original-docx-bytes"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]


def test_compare_rejects_nonpositive_original_page_count(client):
    """original_page_count 必须 >= 1(0 / 负数返回 400)。"""
    import io
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    pdf.new_page(width=595, height=842)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    for bad in (0, -1):
        doc_buf.seek(0)
        pdf_buf.seek(0)
        r = client.post(
            "/api/v1/compare",
            files={
                "source": ("contract.docx", io.BytesIO(doc_buf.getvalue()),
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                "target": ("scan.pdf", io.BytesIO(pdf_buf.getvalue()), "application/pdf"),
            },
            data={"options": f'{{"truncate_to_original_pages": true, "original_page_count": {bad}}}'},
        )
        assert r.status_code == 400, (bad, r.json())


def test_raw_compare_accepts_valid_files(client):
    """无标注版 POST /api/v1/raw-compare 应接受 docx+pdf 并返回 task_id。"""
    import io
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 72), "第一条 测试条款", fontsize=12)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/raw-compare",
        files={
            "source": ("contract.docx", doc_buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", pdf_buf, "application/pdf"),
        },
    )
    assert r.status_code == 200
    assert "task_id" in r.json()


def test_raw_compare_rejects_non_docx_non_pdf(client):
    """无标注版 source 非 docx/pdf 应返回 400。"""
    import io

    r = client.post(
        "/api/v1/raw-compare",
        files={
            "source": ("bad.txt", io.BytesIO(b"text"), "text/plain"),
            "target": ("ok.pdf", io.BytesIO(b"%PDF"), "application/pdf"),
        },
    )
    assert r.status_code == 400
    msg = r.json()["message"]
    assert "docx" in msg and "pdf" in msg


def test_raw_task_not_found_404(client):
    r = client.get("/api/v1/raw-compare/nonexistent")
    assert r.status_code == 404


def test_task_not_found_404(client):
    r = client.get("/api/v1/compare/nonexistent")
    assert r.status_code == 404


def test_uniform_error_has_request_id(client):
    r = client.post("/api/v1/compare", files={})
    body = r.json()
    assert "code" in body
    assert "message" in body
    assert "request_id" in body


# —— PDF 页数上限预检(max_pdf_pages)——

def _make_pdf_bytes(num_pages: int) -> bytes:
    """构造一个 num_pages 页的空白 PDF(供页数上限测试)。"""
    import io

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    buf = io.BytesIO()
    pdf = fitz.open()
    for _ in range(num_pages):
        pdf.new_page(width=595, height=842)
    pdf.save(buf)
    pdf.close()
    buf.seek(0)
    return buf.read()


def _make_docx_bytes() -> bytes:
    import io
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def test_compare_rejects_oversized_pdf(client, monkeypatch):
    """max_pdf_pages 配置后,超过页数的 PDF 直接返回 400 暂不支持。"""
    import io

    monkeypatch.setattr(settings, "max_pdf_pages", 2)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", io.BytesIO(_make_docx_bytes()),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", io.BytesIO(_make_pdf_bytes(3)), "application/pdf"),
        },
    )
    assert r.status_code == 400
    assert "暂不支持" in r.json()["message"]


def test_compare_allows_pdf_within_limit(client, monkeypatch):
    """max_pdf_pages 配置后,页数 <= 上限的 PDF 正常接受。"""
    import io

    monkeypatch.setattr(settings, "max_pdf_pages", 5)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", io.BytesIO(_make_docx_bytes()),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", io.BytesIO(_make_pdf_bytes(3)), "application/pdf"),
        },
    )
    assert r.status_code == 200
    assert "task_id" in r.json()


def test_raw_compare_rejects_oversized_pdf(client, monkeypatch):
    """无标注版同样受 max_pdf_pages 约束。"""
    import io

    monkeypatch.setattr(settings, "max_pdf_pages", 1)

    r = client.post(
        "/api/v1/raw-compare",
        files={
            "source": ("contract.docx", io.BytesIO(_make_docx_bytes()),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", io.BytesIO(_make_pdf_bytes(2)), "application/pdf"),
        },
    )
    assert r.status_code == 400
    assert "暂不支持" in r.json()["message"]


def test_max_pdf_pages_zero_means_unlimited(client, monkeypatch):
    """max_pdf_pages = 0 时不限制,大 PDF 也能通过(默认行为)。"""
    import io

    monkeypatch.setattr(settings, "max_pdf_pages", 0)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", io.BytesIO(_make_docx_bytes()),
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", io.BytesIO(_make_pdf_bytes(20)), "application/pdf"),
        },
    )
    assert r.status_code == 200


def test_llm_config_persists_max_pdf_pages(client, monkeypatch):
    """PUT /api/v1/config/llm 应把 max_pdf_pages 写入 PG llm_config 表并应用到 settings。"""
    from document_comparison.db import repository as db_repo

    # save_llm_overrides 内部合并已有 PG 记录,这里先清空避免被前序用例污染
    db_repo.save_llm_config({})

    r = client.put("/api/v1/config/llm", json={"max_pdf_pages": 12})
    assert r.status_code == 200
    assert r.json()["config"]["max_pdf_pages"] == 12

    # 应用到运行时单例(app.py 复用同一个 settings 对象)
    assert settings.max_pdf_pages == 12

    # 持久化到 PG
    persisted = db_repo.get_llm_config()
    assert persisted["max_pdf_pages"] == 12


def test_llm_config_rejects_negative_max_pdf_pages(client):
    """max_pdf_pages 不能为负数。"""
    r = client.put("/api/v1/config/llm", json={"max_pdf_pages": -1})
    assert r.status_code == 400


def test_llm_config_persists_seal_recovery_settings(client):
    from document_comparison.db import repository as db_repo

    db_repo.save_llm_config({})
    response = client.put(
        "/api/v1/config/llm",
        json={"seal_recovery_enabled": True, "seal_recovery_dpi": 320},
    )

    assert response.status_code == 200
    assert response.json()["config"]["seal_recovery_enabled"] is True
    assert response.json()["config"]["seal_recovery_dpi"] == 320
    assert settings.seal_recovery_enabled is True
    assert settings.seal_recovery_dpi == 320
    persisted = db_repo.get_llm_config()
    assert persisted["seal_recovery_enabled"] is True
    assert persisted["seal_recovery_dpi"] == 320


@pytest.mark.parametrize(
    "payload",
    [
        {"seal_recovery_enabled": "true"},
        {"seal_recovery_dpi": 199},
        {"seal_recovery_dpi": 601},
    ],
)
def test_llm_config_rejects_invalid_seal_recovery_settings(client, payload):
    response = client.put("/api/v1/config/llm", json=payload)
    assert response.status_code == 400


def test_llm_config_persists_llm_direct_diff_prompt(client, monkeypatch):
    """PUT 应把 llm_direct_diff_prompt 写入 PG、应用到 settings 并在 GET 中回读一致。"""
    from document_comparison.db import repository as db_repo

    db_repo.save_llm_config({})
    # 防御性:确保起始态为空(回退内置默认)
    monkeypatch.setattr(settings, "llm_direct_diff_prompt", "")

    custom = "你是合同比对助手。严格输出 JSON:{\"hunks\":[],\"similarity\":1.0}"

    r = client.put("/api/v1/config/llm", json={"llm_direct_diff_prompt": custom})
    assert r.status_code == 200
    assert r.json()["config"]["llm_direct_diff_prompt"] == custom

    # 应用到运行时单例
    assert settings.llm_direct_diff_prompt == custom

    # 持久化到 PG(原值,不脱敏——提示词不是密钥)
    persisted = db_repo.get_llm_config()
    assert persisted["llm_direct_diff_prompt"] == custom

    # GET 回读
    fetched = client.get("/api/v1/config/llm").json()
    assert fetched["llm_direct_diff_prompt"] == custom

    # 空串 PUT → 回退内置默认(settings 字段被清空)
    r2 = client.put("/api/v1/config/llm", json={"llm_direct_diff_prompt": ""})
    assert r2.status_code == 200
    assert r2.json()["config"]["llm_direct_diff_prompt"] == ""
    assert settings.llm_direct_diff_prompt == ""


def test_llm_config_rejects_non_string_llm_direct_diff_prompt(client):
    """llm_direct_diff_prompt 必须是字符串。"""
    r = client.put("/api/v1/config/llm", json={"llm_direct_diff_prompt": 123})
    assert r.status_code == 400
    assert "llm_direct_diff_prompt" in r.json()["message"]


def test_llm_config_exposes_readonly_default_prompt(client):
    """GET/PUT 响应应返回内置默认提示词全文(只读),且 PUT 回传该字段应被拒。"""
    from document_comparison.compare.llm_diff import DEFAULT_DIFF_SYSTEM_PROMPT

    # GET 返回默认提示词全文
    fetched = client.get("/api/v1/config/llm").json()
    assert fetched["llm_direct_diff_default_prompt"] == DEFAULT_DIFF_SYSTEM_PROMPT
    assert "你是合同关键差异比对助手" in fetched["llm_direct_diff_default_prompt"]

    # PUT 响应也带该字段(与 GET 同源)
    r = client.put("/api/v1/config/llm", json={"judge_timeout": settings.judge_timeout})
    assert r.status_code == 200
    assert r.json()["config"]["llm_direct_diff_default_prompt"] == DEFAULT_DIFF_SYSTEM_PROMPT

    # 只读:前端误回传应被白名单挡掉(不写入 settings)
    r2 = client.put(
        "/api/v1/config/llm",
        json={"llm_direct_diff_default_prompt": "试图覆盖默认"},
    )
    assert r2.status_code == 400
    assert "llm_direct_diff_default_prompt" in r2.json()["message"]


def test_llm_config_persists_paddleocr_api_mode(client):
    from document_comparison.db import repository as db_repo

    db_repo.save_llm_config({})
    response = client.put(
        "/api/v1/config/llm",
        json={
            "paddleocr_api_mode": "official_sdk",
            "paddleocr_official_access_token": "official-token-1234",
            "paddleocr_official_model": "PaddleOCR-VL-1.6",
        },
    )

    assert response.status_code == 200
    config = response.json()["config"]
    assert config["paddleocr_api_mode"] == "official_sdk"
    assert config["paddleocr_official_access_token_set"] is True
    assert config["paddleocr_official_access_token"].endswith("1234")
    assert config["paddleocr_official_access_token"] != "official-token-1234"
    persisted = db_repo.get_llm_config()
    assert persisted["paddleocr_api_mode"] == "official_sdk"
    assert persisted["paddleocr_official_access_token"] == "official-token-1234"

    fetched = client.get("/api/v1/config/llm").json()
    assert fetched["persisted"]["paddleocr_official_access_token"] != (
        "official-token-1234"
    )


def test_llm_config_rejects_unknown_paddleocr_api_mode(client):
    response = client.put(
        "/api/v1/config/llm", json={"paddleocr_api_mode": "unknown"}
    )

    assert response.status_code == 400
    assert "paddleocr_api_mode" in response.json()["message"]


def test_llm_config_persists_api_protocol(client):
    """接口协议字段持久化到 PG 并立即生效(llm 与 judge 两组独立)。"""
    from document_comparison.config import settings
    from document_comparison.db import repository as db_repo

    db_repo.save_llm_config({})
    response = client.put(
        "/api/v1/config/llm",
        json={
            "llm_api_protocol": "anthropic",
            "judge_api_protocol": "openai_responses",
        },
    )

    assert response.status_code == 200
    config = response.json()["config"]
    assert config["llm_api_protocol"] == "anthropic"
    assert config["judge_api_protocol"] == "openai_responses"
    assert settings.llm_api_protocol == "anthropic"
    assert settings.judge_api_protocol == "openai_responses"
    persisted = db_repo.get_llm_config()
    assert persisted["llm_api_protocol"] == "anthropic"
    assert persisted["judge_api_protocol"] == "openai_responses"

    fetched = client.get("/api/v1/config/llm").json()
    assert fetched["llm_api_protocol"] == "anthropic"
    assert fetched["judge_api_protocol"] == "openai_responses"


def test_llm_config_rejects_unknown_api_protocol(client):
    response = client.put(
        "/api/v1/config/llm", json={"llm_api_protocol": "rest"}
    )
    assert response.status_code == 400
    assert "llm_api_protocol" in response.json()["message"]

    response = client.put(
        "/api/v1/config/llm", json={"judge_api_protocol": "grpc"}
    )
    assert response.status_code == 400
    assert "judge_api_protocol" in response.json()["message"]


def test_llm_config_rejects_unknown_paddleocr_official_model(client):
    response = client.put(
        "/api/v1/config/llm",
        json={"paddleocr_official_model": "PaddlePaddle/PaddleOCR-VL-1.6"},
    )

    assert response.status_code == 400
    assert "paddleocr_official_model" in response.json()["message"]


def test_external_api_config_persists_applies_and_masks_secret(client):
    """管理端保存外部 API 配置后应立即生效，GET/PUT 均不得返回 Key 明文。"""
    from document_comparison.db import repository as db_repo

    db_repo.save_llm_config({})
    response = client.put(
        "/api/v1/config/llm",
        json={
            "external_api_key": "external-secret-1234",
            "external_public_base_url": "https://compare.example.com/",
            "external_max_upload_mb": 80,
            "external_image_dpi": 180,
            "external_ocr_backend": "llm",
            "external_enable_llm_judge": True,
            "external_enable_llm_alignment": True,
            "external_enable_risk_assessment": True,
            "external_truncate_to_original_pages": True,
        },
    )
    assert response.status_code == 200, response.json()
    config = response.json()["config"]
    assert config["external_api_key_set"] is True
    assert config["external_api_key"].endswith("1234")
    assert config["external_api_key"] != "external-secret-1234"
    assert config["external_public_base_url"] == "https://compare.example.com"
    assert config["external_max_upload_mb"] == 80
    assert config["external_image_dpi"] == 180
    assert config["external_ocr_backend"] == "llm"
    assert config["external_enable_llm_judge"] is True
    assert config["external_enable_llm_alignment"] is True
    assert config["external_enable_risk_assessment"] is True
    assert config["external_truncate_to_original_pages"] is True
    assert config["external_enabled"] is True

    assert settings.external_api_key == "external-secret-1234"
    assert settings.external_public_base_url == "https://compare.example.com"
    assert settings.external_truncate_to_original_pages is True
    persisted = db_repo.get_llm_config()
    assert persisted["external_api_key"] == "external-secret-1234"

    fetched = client.get("/api/v1/config/llm")
    assert fetched.status_code == 200
    body = fetched.json()
    assert body["external_truncate_to_original_pages"] is True
    assert body["external_api_key"] != "external-secret-1234"
    assert body["persisted"]["external_api_key"] != "external-secret-1234"
    # 现有模型 Key 的持久化快照也一并保持脱敏。
    assert body["persisted"].get("llm_api_key", "") != settings.llm_api_key or not settings.llm_api_key


@pytest.mark.parametrize(
    "payload",
    [
        {"external_public_base_url": "not-a-url"},
        {"external_public_base_url": "https://user:pass@example.com"},
        {"external_max_upload_mb": 0},
        {"external_max_upload_mb": 1025},
        {"external_image_dpi": 71},
        {"external_image_dpi": 601},
        {"external_ocr_backend": "other"},
        {"external_enable_llm_judge": "true"},
        {"external_enable_llm_alignment": "true"},
        {"external_enable_risk_assessment": 1},
        {"external_truncate_to_original_pages": "yes"},
    ],
)
def test_external_api_config_rejects_invalid_values(client, payload):
    response = client.put("/api/v1/config/llm", json=payload)
    assert response.status_code == 400


# —— 金额统计端点冒烟(/api/v1/statement)——

def test_statement_accepts_multiple_pdfs(client, monkeypatch):
    """POST /api/v1/statement 接受多个 PDF,返回 task_id + file_count。

    mock 掉 run_statement 避免触发真实 OCR;TestClient 多文件需用 list-of-tuple 形式
    (同字段名多次出现),不能用 dict 单键 list 形式(后者会触发 multipart 解析错误)。
    """
    import io

    from document_comparison.api.app import task_manager

    async def _noop(*args, **kwargs):
        return None
    monkeypatch.setattr(task_manager, "run_statement", _noop)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    def _pdf_bytes():
        buf = io.BytesIO()
        pdf = fitz.open()
        pdf.new_page(width=595, height=842)
        pdf.save(buf)
        pdf.close()
        buf.seek(0)
        return buf

    r = client.post(
        "/api/v1/statement",
        files=[
            ("target", ("a.pdf", _pdf_bytes(), "application/pdf")),
            ("target", ("b.pdf", _pdf_bytes(), "application/pdf")),
        ],
    )
    assert r.status_code == 200
    body = r.json()
    assert "task_id" in body
    assert body["file_count"] == 2


def test_statement_rejects_non_pdf(client):
    """混入非 PDF → 400。"""
    import io

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    pdf.new_page(width=595, height=842)
    pdf.save(pdf_buf)
    pdf.close()
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/statement",
        files=[
            ("target", ("ok.pdf", pdf_buf, "application/pdf")),
            ("target", ("bad.txt", io.BytesIO(b"text"), "text/plain")),
        ],
    )
    assert r.status_code == 400
    assert "pdf" in r.json()["message"].lower()


def test_statement_rejects_oversized_pdf(client, monkeypatch):
    """多文件中任一超 max_pdf_pages → 400 暂不支持。"""
    import io

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore

    monkeypatch.setattr(settings, "max_pdf_pages", 1)
    # 构造 3 页 PDF
    buf = io.BytesIO()
    pdf = fitz.open()
    for _ in range(3):
        pdf.new_page(width=595, height=842)
    pdf.save(buf)
    pdf.close()
    buf.seek(0)

    r = client.post(
        "/api/v1/statement",
        files=[
            ("target", ("big.pdf", buf, "application/pdf")),
        ],
    )
    assert r.status_code == 400
    assert "暂不支持" in r.json()["message"]


def test_statement_task_not_found_404(client):
    r = client.get("/api/v1/statement/nonexistent")
    assert r.status_code == 404


def test_statement_report_not_found_404(client):
    r = client.get("/api/v1/statement/nonexistent/report")
    assert r.status_code == 404


# —— 比对记录历史端点 + 进程重启兜底测试(/api/v1/tasks)——

def test_tasks_listing_returns_items(client):
    """POST 提交后会写入 task_records,GET /api/v1/tasks 应能列出。"""
    import io

    from document_comparison.api.app import task_manager

    # mock run 避免真实 OCR
    async def _noop(*args, **kwargs):
        return None
    orig_run = task_manager.run
    task_manager.run = _noop  # type: ignore[assignment]

    try:
        import pymupdf as fitz
        pdf_buf = io.BytesIO()
        pdf = fitz.open()
        pdf.new_page(width=595, height=842)
        pdf.save(pdf_buf)
        pdf.close()
        pdf_buf.seek(0)

        from docx import Document  # type: ignore[import-untyped]
        doc = Document()
        doc.add_paragraph("第一条 测试条款")
        doc_buf = io.BytesIO()
        doc.save(doc_buf)
        doc_buf.seek(0)

        r = client.post(
            "/api/v1/compare",
            files={
                "source": ("contract.docx", doc_buf,
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                "target": ("scan.pdf", pdf_buf, "application/pdf"),
            },
        )
        assert r.status_code == 200
    finally:
        task_manager.run = orig_run  # type: ignore[assignment]

    # 列表端点
    r = client.get("/api/v1/tasks")
    assert r.status_code == 200
    body = r.json()
    assert "items" in body and "total" in body
    assert body["total"] >= 1
    assert any(it["kind"] == "compare" for it in body["items"])


def test_tasks_listing_filter_by_kind(client):
    """kind 过滤生效:仅 compare。"""
    from document_comparison.db import repository as db_repo

    db_repo.create_task("list-c1", "compare", target_names=["t.pdf"])
    db_repo.create_task("list-r1", "raw", target_names=["t.pdf"])
    db_repo.create_task("list-s1", "statement", target_names=["t.pdf"])

    r = client.get("/api/v1/tasks?kind=compare")
    assert r.status_code == 200
    body = r.json()
    assert all(it["kind"] == "compare" for it in body["items"])
    assert any(it["task_id"] == "list-c1" for it in body["items"])
    assert all(it["task_id"] != "list-r1" for it in body["items"])


def test_tasks_listing_rejects_bad_kind(client):
    r = client.get("/api/v1/tasks?kind=unknown")
    assert r.status_code == 400


def test_tasks_listing_filter_by_document_type(client):
    """document_type 过滤(1=发票 / 2=对帐单)生效,且序列化带该字段。"""
    from document_comparison.db import repository as db_repo

    db_repo.create_task(
        "list-sc1", "statement", target_names=["t.pdf"], document_type="1"
    )
    db_repo.create_task(
        "list-ss1", "statement", target_names=["t.pdf"], document_type="2"
    )
    db_repo.create_task("list-sn1", "statement", target_names=["t.pdf"])
    db_repo.create_task("list-cn1", "compare", target_names=["t.pdf"])

    r = client.get("/api/v1/tasks?document_type=1")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] >= 1
    assert all(it["document_type"] == "1" for it in body["items"])
    assert any(it["task_id"] == "list-sc1" for it in body["items"])

    r = client.get("/api/v1/tasks?document_type=2")
    assert r.status_code == 200
    ids = {it["task_id"] for it in r.json()["items"]}
    assert "list-ss1" in ids
    assert "list-sc1" not in ids  # 发票任务不混入

    # 序列化字段:compare 任务与未设置的 statement 任务为 None
    r = client.get("/api/v1/tasks")
    items = {it["task_id"]: it for it in r.json()["items"]}
    assert items["list-cn1"]["document_type"] is None
    assert items["list-sn1"]["document_type"] is None
    assert items["list-sc1"]["document_type"] == "1"


def test_tasks_listing_rejects_bad_document_type(client):
    r = client.get("/api/v1/tasks?document_type=3")
    assert r.status_code == 400


def test_get_compare_result_survives_no_inmemory_task(client):
    """进程重启模拟:清空内存 task_manager._tasks 后,GET 应从 PG 回兜返回 done 状态。"""
    from document_comparison.db import repository as db_repo
    from document_comparison.models import TamperReport
    from document_comparison.api.app import task_manager

    task_manager._tasks.clear()

    db_repo.create_task("restart-1", "compare", target_names=["t.pdf"])
    db_repo.update_task_status(
        "restart-1", "done",
        overall_risk="clean",
        change_status="clean",
        elapsed=1.23,
        finished=True,
    )
    db_repo.save_compare_report(
        "restart-1", TamperReport(source="a.docx", target="b.pdf", overall_risk="clean"),
    )

    r = client.get("/api/v1/compare/restart-1")
    assert r.status_code == 200
    body = r.json()
    assert body["task_id"] == "restart-1"
    assert body["status"] == "done"
    assert body["overall_risk"] == "clean"
    assert body["elapsed"] == pytest.approx(1.23)
    assert body["report"]["overall_risk"] == "clean"


def test_download_report_from_pg_after_inmemory_cleared(client):
    """下载端点在内存 Task 不存在时应从 PG JSONB 还原报告(原文件兜底已移除)。

    回归保护:此前 /report 端点三级兜底「内存→文件→PG」,本次改造去掉文件层,
    仅靠 PG 必须仍能返回完整 JSON。
    """
    from document_comparison.db import repository as db_repo
    from document_comparison.models import TamperReport
    from document_comparison.api.app import task_manager

    task_manager._tasks.clear()

    db_repo.create_task("dl-1", "compare", target_names=["t.pdf"])
    db_repo.update_task_status("dl-1", "done", finished=True)
    db_repo.save_compare_report(
        "dl-1",
        TamperReport(source="a.docx", target="b.pdf", overall_risk="clean"),
    )

    r = client.get("/api/v1/compare/dl-1/report?format=json")
    assert r.status_code == 200
    body = r.json()
    assert body["source"] == "a.docx"
    assert body["target"] == "b.pdf"
    assert body["overall_risk"] == "clean"


def test_get_task_events_endpoint(client):
    """/api/v1/tasks/{id}/events 返回持久化的里程碑事件。"""
    from document_comparison.db import repository as db_repo

    db_repo.create_task("ev-1", "compare")
    db_repo.save_milestone_event("ev-1", "start", 0.0, {})
    db_repo.save_milestone_event("ev-1", "done", 1.0, {"word": 0.3})

    r = client.get("/api/v1/tasks/ev-1/events")
    assert r.status_code == 200
    body = r.json()
    assert body["task_id"] == "ev-1"
    assert [it["stage"] for it in body["items"]] == ["start", "done"]


def test_get_task_events_unknown_404(client):
    r = client.get("/api/v1/tasks/unknown-task/events")
    assert r.status_code == 404


# —— LLM 调用记录端点(/api/v1/tasks/{task_id}/llm-calls)——

def test_get_task_llm_calls_endpoint(client):
    """/api/v1/tasks/{id}/llm-calls 返回持久化的 LLM 调用记录(按 id 升序)。"""
    from document_comparison.observability import LlmCallRecord
    from document_comparison.db import repository as db_repo

    db_repo.create_task("lc-1", "compare")
    db_repo.save_llm_calls_batch("lc-1", [
        LlmCallRecord(
            kind="ocr", attempt=1, payload={"model": "m"},
            status_code=200, elapsed_ms=50,
            response={"choices": [{"message": {"content": "p1"}}]},
        ),
        LlmCallRecord(
            kind="judge", attempt=1, payload={"model": "j"},
            status_code=429, error="transient HTTP 429",
        ),
    ])

    r = client.get("/api/v1/tasks/lc-1/llm-calls")
    assert r.status_code == 200
    body = r.json()
    assert body["task_id"] == "lc-1"
    assert len(body["items"]) == 2
    assert [it["kind"] for it in body["items"]] == ["ocr", "judge"]
    assert body["items"][0]["status_code"] == 200
    assert body["items"][0]["response"]["choices"][0]["message"]["content"] == "p1"
    assert body["items"][1]["status_code"] == 429
    assert body["items"][1]["error"] == "transient HTTP 429"


def test_get_task_llm_calls_unknown_404(client):
    r = client.get("/api/v1/tasks/unknown-task/llm-calls")
    assert r.status_code == 404


def test_get_task_llm_calls_empty_for_task_without_calls(client):
    """任务存在但无 LLM 调用记录 → 空列表(不报错)。"""
    from document_comparison.db import repository as db_repo

    db_repo.create_task("lc-empty", "compare")
    r = client.get("/api/v1/tasks/lc-empty/llm-calls")
    assert r.status_code == 200
    assert r.json()["items"] == []


def test_compare_rejects_429_when_pg_count_at_limit(client, monkeypatch):
    """PG 计数达到 max_concurrent_tasks 时,提交端点返回 429。

    多 worker 全局并发上限走 db_repo.count_active_tasks;此测试直接打桩计数,
    验证限流逻辑(不依赖内存 _tasks,跨 worker 一致)。
    """
    import io

    from docx import Document  # type: ignore[import-untyped]

    from document_comparison.db import repository as db_repo

    # 计数打桩为满值(= max_concurrent_tasks,默认 4)
    monkeypatch.setattr(db_repo, "count_active_tasks", lambda: settings.max_concurrent_tasks)

    doc = Document()
    doc.add_paragraph("第一条 测试条款")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    doc_buf.seek(0)

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    pdf_buf = io.BytesIO()
    pdf = fitz.open()
    pdf.new_page(width=595, height=842)
    pdf.save(pdf_buf)
    pdf_buf.seek(0)

    r = client.post(
        "/api/v1/compare",
        files={
            "source": ("contract.docx", doc_buf,
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "target": ("scan.pdf", pdf_buf, "application/pdf"),
        },
    )
    assert r.status_code == 429
    assert "并发" in r.json()["message"]


# —— 每日调用统计端点(/api/v1/stats/daily)——

def test_daily_stats_endpoint(client):
    """/api/v1/stats/daily 返回按天聚合序列,days 越界(0 或 91)被 422 拒绝。"""
    from document_comparison.db import repository as db_repo

    db_repo.create_task("stats-1", "compare", target_names=["t.pdf"])
    db_repo.update_task_status("stats-1", "done", finished=True)

    r = client.get("/api/v1/stats/daily")
    assert r.status_code == 200
    body = r.json()
    assert body["days"] == 14
    assert len(body["series"]) == 14
    assert [p["date"] for p in body["series"]] == sorted(p["date"] for p in body["series"])
    # 刚创建的任务落在最后一个(北京时间今天)桶
    last = body["series"][-1]
    assert last["tasks"]["total"] == 1
    assert last["tasks"]["by_kind"].get("compare") == 1
    assert last["tasks"]["by_status"].get("done") == 1

    assert client.get("/api/v1/stats/daily?days=0").status_code == 422
    assert client.get("/api/v1/stats/daily?days=91").status_code == 422


def test_daily_stats_custom_range(client):
    """/api/v1/stats/daily?start=&end= 自由区间;非法格式/区间/跨度返回 400。"""
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    from document_comparison.db import repository as db_repo

    beijing = ZoneInfo("Asia/Shanghai")
    today = datetime.now(beijing).date()
    db_repo.create_task("csr-1", "compare", target_names=["t.pdf"])

    r = client.get(
        f"/api/v1/stats/daily?start={(today - timedelta(days=2)).isoformat()}"
        f"&end={today.isoformat()}"
    )
    assert r.status_code == 200
    body = r.json()
    assert body["days"] == 3
    assert [p["date"] for p in body["series"]] == [
        (today - timedelta(days=i)).isoformat() for i in (2, 1, 0)
    ]
    assert body["series"][-1]["tasks"]["total"] == 1

    # 未来 end 截断到今天(浏览器本地时区晚于北京时的容错)
    r = client.get(f"/api/v1/stats/daily?end={(today + timedelta(days=3)).isoformat()}")
    assert r.status_code == 200
    assert r.json()["series"][-1]["date"] == today.isoformat()

    assert client.get("/api/v1/stats/daily?start=2026/08/01").status_code == 400
    assert client.get(
        f"/api/v1/stats/daily?start={today.isoformat()}&end={(today - timedelta(days=1)).isoformat()}"
    ).status_code == 400
    assert client.get(
        f"/api/v1/stats/daily?start={(today - timedelta(days=400)).isoformat()}"
    ).status_code == 400
